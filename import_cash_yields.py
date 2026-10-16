"""Convert a local Korean government-bond yield workbook to the risk CSV.

The input workbook stays unchanged. Run with a Python environment that has
openpyxl installed; the dashboard itself only reads the resulting CSV.
"""
from __future__ import annotations

import argparse
import csv
import math
from datetime import date, datetime
from pathlib import Path


def convert(source: Path, destination: Path) -> tuple[int, date, date]:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        if "일자" not in headers or "국고채(3년)" not in headers or "국고채(10년)" not in headers:
            raise ValueError("Expected columns: 일자, 국고채(3년), 국고채(10년)")
        columns = [headers.index(name) for name in ("일자", "국고채(3년)", "국고채(10년)")]
        normalized = []
        for row_number, row in enumerate(rows, 2):
            values = [row[index] if index < len(row) else None for index in columns]
            if all(value is None for value in values):
                continue
            day = values[0].date() if isinstance(values[0], datetime) else values[0]
            if not isinstance(day, date) or any(isinstance(value, bool) or
                                                not isinstance(value, (int, float)) or
                                                not math.isfinite(value) for value in values[1:]):
                raise ValueError(f"Invalid date or yields at row {row_number}")
            normalized.append((day, float(values[1]), float(values[2])))
    finally:
        workbook.close()
    normalized.sort(key=lambda item: item[0])
    if not normalized or any(current[0] == before[0] for before, current in
                             zip(normalized, normalized[1:])):
        raise ValueError("Input needs at least one dated row with no duplicate dates")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(("date", "cash_3y_yield_pct", "cash_10y_yield_pct"))
        writer.writerows((day.isoformat(), three, ten) for day, three, ten in normalized)
    return len(normalized), normalized[0][0], normalized[-1][0]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    count, first, last = convert(args.source, args.destination)
    print(f"Converted {count} daily records: {first} to {last}")


if __name__ == "__main__":
    main()
